import hashlib
import hmac
import json
import os
import tempfile
import unittest
from io import BytesIO
from unittest.mock import patch

os.environ.setdefault("FLASK_DEBUG", "1")
os.environ.setdefault("BOOTSTRAP_ADMIN_PASSWORD", "TestBootstrap@123")
os.environ["DATABASE_URL"] = "sqlite:///" + tempfile.mktemp(suffix=".db").replace("\\", "/")

from openpyxl import load_workbook
from werkzeug.security import generate_password_hash

import run


class DashboardTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.context = run.app.app_context()
        cls.context.push()
        run.db.create_all()
        cls.school = run.School(name="Dashboard Test School", academic_year="2026/2027", term="Term 1", onboarded=True)
        run.db.session.add(cls.school)
        run.db.session.flush()
        cls.users = {}
        for role in [
            "system_admin",
            "school_admin",
            "teacher",
            "student",
            "accountant",
            "registrar",
            "receptionist",
            "librarian",
        ]:
            account = run.User(
                school_id=None if role == "system_admin" else cls.school.id,
                role=role,
                full_name=role.replace("_", " ").title(),
                username=f"dashboard_{role}",
                password_hash=generate_password_hash("Test@12345"),
                must_change_password=False,
            )
            run.db.session.add(account)
            cls.users[role] = account
        run.db.session.commit()
        student_account = cls.users["student"]
        run.db.session.add(run.Student(school_id=cls.school.id, user_id=student_account.id, admission_no="TEST001"))
        run.db.session.commit()
        cls.student = run.Student.query.filter_by(user_id=student_account.id).first()
        cls.parent = run.User(
            school_id=cls.school.id,
            role="parent",
            full_name="Test Parent",
            username="dashboard_parent",
            email="parent@example.com",
            password_hash=generate_password_hash("Test@12345"),
            must_change_password=False,
        )
        run.db.session.add(cls.parent)
        run.db.session.flush()
        run.db.session.add(
            run.ParentStudent(school_id=cls.school.id, parent_id=cls.parent.id, student_id=cls.student.id)
        )
        run.db.session.commit()
        run.Config.PAYSTACK_SECRET_KEY = "sk_test_for_unit_tests_only"
        run.Config.PAYSTACK_PUBLIC_KEY = "pk_test_for_unit_tests_only"
        run.Config.PARENT_REPORT_FEE = "10.00"
        run.Config.PAYSTACK_CURRENCY = "GHS"

    @classmethod
    def tearDownClass(cls):
        run.db.session.remove()
        cls.context.pop()

    def test_every_staff_dashboard_renders_reference_layout(self):
        client = run.app.test_client()
        for role, account in self.users.items():
            with self.subTest(role=role):
                with client.session_transaction() as session:
                    session["user_id"] = account.id
                    session["session_version"] = account.session_version
                    session["_csrf"] = "test-csrf"
                response = client.get("/dashboard")
                self.assertEqual(response.status_code, 200)
                html = response.get_data(as_text=True)
                self.assertIn("kpi-grid", html)
                if role == "student":
                    self.assertIn("Recent Results", html)
                    self.assertIn("Fee Balance", html)
                    self.assertIn("Quick Links", html)
                else:
                    self.assertIn("Quick Actions", html)
                    self.assertIn("line-chart", html)
                self.assertIn("theme-toggle", html)
                self.assertIn("smart-school-theme", html)

    def test_logo_is_packaged(self):
        response = run.app.test_client().get("/static/smart-school-logo.png")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "image/png")

    def test_landing_page_matches_product_direction(self):
        html = run.app.test_client().get("/").get_data(as_text=True)
        self.assertIn("Smarter School", html)
        self.assertIn("Better Future", html)
        self.assertIn("Student Management", html)
        self.assertIn("product-window", html)
        self.assertIn("About Us", html)
        self.assertIn("Get Started", html)
        for damaged_character in ("â", "Â", "ï¼"):
            self.assertNotIn(damaged_character, html)

    def test_login_uses_reference_role_tabs_and_accessible_fields(self):
        html = run.app.test_client().get("/login?portal=admin").get_data(as_text=True)
        self.assertIn("auth-brand-panel", html)
        self.assertIn("Administrator", html)
        self.assertIn("Teacher", html)
        self.assertIn("Parent", html)
        self.assertIn("Student", html)
        self.assertIn("Email address or username", html)
        self.assertIn("Remember me", html)
        self.assertIn("Forgot password?", html)
        self.assertIn("Enter your username", html)
        self.assertIn("auth-brand-panel.role-admin", html)
        self.assertIn("auth-brand-panel.role-teacher", html)
        self.assertIn("auth-brand-panel.role-parent", html)
        self.assertIn("auth-brand-panel.role-student", html)

    def test_login_accepts_username_not_email(self):
        account = self.users["school_admin"]
        account.email = "school-admin-login@example.com"
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        rejected = client.post(
            "/login?portal=admin",
            data={
                "_csrf": "test-csrf",
                "portal": "admin",
                "username": account.email,
                "password": "Test@12345",
            },
        )
        self.assertEqual(rejected.status_code, 200)
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        accepted = client.post(
            "/login?portal=admin",
            data={
                "_csrf": "test-csrf",
                "portal": "admin",
                "username": account.username,
                "password": "Test@12345",
            },
        )
        self.assertEqual(accepted.status_code, 302)
        self.assertIn("/dashboard", accepted.location)

    def test_login_checks_the_correct_school_when_usernames_repeat(self):
        first_school = run.School(name="First Duplicate Username School")
        second_school = run.School(name="Second Duplicate Username School")
        run.db.session.add_all([first_school, second_school])
        run.db.session.flush()
        run.db.session.add_all(
            [
                run.User(
                    school_id=first_school.id,
                    role="school_admin",
                    full_name="First Ezra",
                    username="ezra",
                    password_hash=generate_password_hash("Wrong@12345"),
                    must_change_password=False,
                ),
                run.User(
                    school_id=second_school.id,
                    role="school_admin",
                    full_name="Second Ezra",
                    username="ezra",
                    password_hash=generate_password_hash("Correct@12345"),
                    must_change_password=False,
                ),
            ]
        )
        run.db.session.commit()

        client = run.app.test_client()
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        response = client.post(
            "/login?portal=admin",
            data={
                "_csrf": "test-csrf",
                "portal": "admin",
                "username": "Ezra",
                "password": "Correct@12345",
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertIn("/dashboard", response.location)
        with client.session_transaction() as session:
            logged_in = run.db.session.get(run.User, session["user_id"])
        self.assertEqual(logged_in.full_name, "Second Ezra")

    def test_school_admin_dashboard_exposes_import_and_export(self):
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.users["school_admin"].id
            session["session_version"] = self.users["school_admin"].session_version
            session["_csrf"] = "test-csrf"
        html = client.get("/dashboard").get_data(as_text=True)
        self.assertIn("Import Excel", html)
        self.assertIn("Export Excel", html)
        self.assertIn("/admin/export.xlsx", html)

    def test_school_admin_export_contains_school_records(self):
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.users["school_admin"].id
            session["session_version"] = self.users["school_admin"].session_version
            session["_csrf"] = "test-csrf"
        response = client.get("/admin/export.xlsx")
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.data), read_only=True)
        self.assertEqual(workbook.sheetnames, ["Students", "Teachers"])
        student_rows = list(workbook["Students"].iter_rows(values_only=True))
        teacher_rows = list(workbook["Teachers"].iter_rows(values_only=True))
        self.assertEqual(student_rows[0][0:4], ("full_name", "username", "admission_no", "class"))
        self.assertTrue(any(row[1] == "dashboard_student" for row in student_rows[1:]))
        self.assertTrue(any(row[1] == "dashboard_teacher" for row in teacher_rows[1:]))

    def test_students_page_renders_for_admin_and_teacher(self):
        client = run.app.test_client()
        for role in ["school_admin", "teacher"]:
            with self.subTest(role=role):
                with client.session_transaction() as session:
                    session["user_id"] = self.users[role].id
                    session["session_version"] = self.users[role].session_version
                    session["_csrf"] = "test-csrf"
                response = client.get("/students")
                self.assertEqual(response.status_code, 200, response.get_data(as_text=True))

    def test_dashboard_search_finds_students_and_staff(self):
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.users["school_admin"].id
            session["session_version"] = self.users["school_admin"].session_version
            session["_csrf"] = "test-csrf"
        student_results = client.get("/search?q=Student")
        self.assertEqual(student_results.status_code, 200)
        self.assertIn("TEST001", student_results.get_data(as_text=True))
        staff_results = client.get("/search?q=Teacher")
        self.assertEqual(staff_results.status_code, 200)
        self.assertIn("dashboard_teacher", staff_results.get_data(as_text=True))

    def test_report_includes_class_and_promotion_status(self):
        student = run.Student.query.filter_by(user_id=self.users["student"].id).first()
        class_group = run.ClassRoom(school_id=self.school.id, name="JHS 2")
        run.db.session.add(class_group)
        run.db.session.flush()
        student.class_id = class_group.id
        student.promotion_note = "Promoted from JHS 1 to JHS 2"
        subject = run.Subject(school_id=self.school.id, name="Integrated Science", code="SCI")
        run.db.session.add(subject)
        run.db.session.flush()
        run.db.session.add(
            run.Score(
                school_id=self.school.id,
                student_id=student.id,
                subject_id=subject.id,
                class_score=35,
                exam_score=50,
                term=self.school.term,
                academic_year=self.school.academic_year,
                workflow_status="published",
                published_at=run.datetime.utcnow(),
            )
        )
        run.db.session.commit()
        Payment = run.app.feature_models["StudentReportPayment"]
        run.db.session.add(
            Payment(
                school_id=self.school.id,
                user_id=self.users["student"].id,
                student_id=student.id,
                academic_year=self.school.academic_year,
                term=self.school.term,
                reference="student-report-existing-test",
                amount_subunit=1000,
                currency="GHS",
                status="success",
            )
        )
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.users["student"].id
            session["session_version"] = self.users["student"].session_version
            session["_csrf"] = "test-csrf"
        html = client.get("/my-results").get_data(as_text=True)
        self.assertIn("JHS 2", html)
        self.assertIn("Promoted from JHS 1 to JHS 2", html)
        self.assertIn("Terminal Report", html)
        self.assertIn("Class Score", html)
        self.assertIn("Head of School", html)
        pdf = client.get("/my-results.pdf")
        self.assertEqual(pdf.status_code, 200)
        self.assertEqual(pdf.mimetype, "application/pdf")
        self.assertTrue(pdf.data.startswith(b"%PDF"))

    def test_staff_report_preview_and_publish_controls(self):
        student = run.Student.query.filter_by(user_id=self.users["student"].id).first()
        if not student.class_id:
            class_group = run.ClassRoom(school_id=self.school.id, name="Preview Class")
            run.db.session.add(class_group)
            run.db.session.flush()
            student.class_id = class_group.id
            run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.users["school_admin"].id
            session["session_version"] = self.users["school_admin"].session_version
            session["_csrf"] = "test-csrf"
        listing = client.get("/report-cards")
        self.assertEqual(listing.status_code, 200)
        preview = client.get(f"/report-cards/{student.id}/preview")
        html = preview.get_data(as_text=True)
        self.assertEqual(preview.status_code, 200)
        self.assertIn("Publish Report", html)
        self.assertIn("reference-report", html)
        published = client.post(
            f"/report-cards/{student.id}/publish", data={"_csrf": "test-csrf"}, follow_redirects=True
        )
        self.assertEqual(published.status_code, 200)

    def test_published_score_requires_reason_and_records_correction(self):
        subject = run.Subject(school_id=self.school.id, name="Correction Audit Subject", code="CAS")
        run.db.session.add(subject)
        run.db.session.flush()
        score = run.Score(
            school_id=self.school.id,
            student_id=self.student.id,
            subject_id=subject.id,
            class_score=20,
            exam_score=50,
            term=self.school.term,
            academic_year=self.school.academic_year,
            workflow_status="published",
            published_at=run.datetime.utcnow(),
            locked_at=run.datetime.utcnow(),
            revision=1,
        )
        run.db.session.add(score)
        run.db.session.commit()
        client = run.app.test_client()
        admin = self.users["school_admin"]
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
            session["_csrf"] = "test-csrf"
        payload = {
            "_csrf": "test-csrf",
            "student_id": self.student.id,
            "subject_id": subject.id,
            "class_score": "25",
            "exam_score": "55",
            "term": self.school.term,
            "academic_year": self.school.academic_year,
        }
        rejected = client.post("/scores", data=payload, follow_redirects=True)
        self.assertIn("correction reason is required", rejected.get_data(as_text=True).lower())
        run.db.session.refresh(score)
        self.assertEqual(score.class_score, 20)

        payload["correction_reason"] = "Corrected verified entry"
        accepted = client.post("/scores", data=payload, follow_redirects=True)
        self.assertIn("Score saved", accepted.get_data(as_text=True))
        run.db.session.refresh(score)
        self.assertEqual(score.workflow_status, "draft")
        self.assertIsNone(score.locked_at)
        self.assertEqual(score.revision, 2)
        ResultChange = run.app.platform_models["ResultChange"]
        change = ResultChange.query.filter_by(score_id=score.id).one()
        self.assertEqual(change.reason, "Corrected verified entry")
        self.assertEqual(change.old_class_score, 20)
        self.assertEqual(change.new_class_score, 25)

    def test_student_report_is_locked_without_payment(self):
        Payment = run.app.feature_models["StudentReportPayment"]
        Payment.query.filter_by(user_id=self.users["student"].id).delete()
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.users["student"].id
            session["session_version"] = self.users["student"].session_version
            session["_csrf"] = "test-csrf"
        response = client.get("/my-results")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/student/report/payment", response.location)
        pdf = client.get("/my-results.pdf")
        self.assertEqual(pdf.status_code, 302)

    def test_admin_forgot_password_uses_generic_response(self):
        self.users["school_admin"].email = "admin-reset@example.com"
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        with patch("smart_features._send_resend") as sender:
            response = client.post(
                "/forgot-password",
                data={"_csrf": "test-csrf", "email": "admin-reset@example.com"},
                follow_redirects=True,
            )
        self.assertEqual(response.status_code, 200)
        self.assertIn("If that administrator email exists", response.get_data(as_text=True))
        sender.assert_called_once()

    def test_parent_report_is_locked_until_verified_payment(self):
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = self.parent.id
            session["session_version"] = self.parent.session_version
            session["_csrf"] = "test-csrf"
        response = client.get(f"/parent/report/{self.student.id}")
        self.assertEqual(response.status_code, 302)
        self.assertIn("/payment", response.location)

        with patch("run.paystack_request") as paystack:
            paystack.return_value = {
                "status": True,
                "data": {
                    "authorization_url": "https://checkout.paystack.com/test-access",
                    "reference": "ignored",
                },
            }
            initialized = client.post(f"/parent/report/{self.student.id}/paystack", data={"_csrf": "test-csrf"})
        self.assertEqual(initialized.status_code, 302)
        self.assertTrue(initialized.location.startswith("https://checkout.paystack.com/"))
        payment = (
            run.ParentReportPayment.query.filter_by(parent_id=self.parent.id, student_id=self.student.id)
            .order_by(run.ParentReportPayment.id.desc())
            .first()
        )
        self.assertEqual(payment.status, "pending")

        with patch("run.paystack_request") as paystack:
            paystack.return_value = {
                "status": True,
                "data": {
                    "status": "success",
                    "reference": payment.reference,
                    "amount": 1000,
                    "currency": "GHS",
                    "id": 12345,
                },
            }
            verified = client.get(f"/payments/paystack/callback?reference={payment.reference}", follow_redirects=True)
        self.assertEqual(verified.status_code, 200)
        self.assertIn("Terminal Report", verified.get_data(as_text=True))
        self.assertEqual(payment.status, "success")

    def test_paystack_webhook_requires_valid_signature(self):
        body = json.dumps({"event": "charge.success", "data": {"reference": "unknown-reference"}}).encode()
        client = run.app.test_client()
        self.assertEqual(
            client.post("/payments/paystack/webhook", data=body, content_type="application/json").status_code, 401
        )
        signature = hmac.new(run.Config.PAYSTACK_SECRET_KEY.encode(), body, hashlib.sha512).hexdigest()
        response = client.post(
            "/payments/paystack/webhook",
            data=body,
            content_type="application/json",
            headers={"x-paystack-signature": signature},
        )
        self.assertEqual(response.status_code, 200)

    def test_health_endpoint_checks_database_without_authentication(self):
        response = run.app.test_client().get("/health")
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload["status"], "ok")
        self.assertEqual(payload["database"], "ok")

    def test_platform_owner_portal_is_restricted(self):
        owner_client = run.app.test_client()
        owner = self.users["system_admin"]
        with owner_client.session_transaction() as session:
            session["user_id"] = owner.id
            session["session_version"] = owner.session_version
            session["_csrf"] = "test-csrf"
        response = owner_client.get("/owner")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Platform Owner", response.get_data(as_text=True))

        school_client = run.app.test_client()
        admin = self.users["school_admin"]
        with school_client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
            session["_csrf"] = "test-csrf"
        self.assertNotEqual(school_client.get("/owner").status_code, 200)

    def test_logout_all_devices_invalidates_existing_session(self):
        client = run.app.test_client()
        admin = self.users["school_admin"]
        original_version = admin.session_version
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = original_version
            session["_csrf"] = "test-csrf"
        response = client.post("/logout-all", data={"_csrf": "test-csrf"})
        self.assertEqual(response.status_code, 302)
        run.db.session.refresh(admin)
        self.assertEqual(admin.session_version, original_version + 1)
        self.assertEqual(client.get("/dashboard").status_code, 302)

    def test_cross_school_report_access_is_denied(self):
        other_school = run.School(
            name="Isolated School",
            slug="isolated-school",
            academic_year=self.school.academic_year,
            term=self.school.term,
            onboarded=True,
        )
        run.db.session.add(other_school)
        run.db.session.flush()
        other_user = run.User(
            school_id=other_school.id,
            role="student",
            full_name="Other Student",
            username="other_student",
            password_hash=generate_password_hash("Other@Test123"),
            must_change_password=False,
        )
        run.db.session.add(other_user)
        run.db.session.flush()
        other_student = run.Student(school_id=other_school.id, user_id=other_user.id, admission_no="OTHER001")
        run.db.session.add(other_student)
        run.db.session.commit()

        client = run.app.test_client()
        admin = self.users["school_admin"]
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
            session["_csrf"] = "test-csrf"
        response = client.get(f"/report-cards/{other_student.id}/preview")
        self.assertEqual(response.status_code, 404)

    def test_suspended_school_blocks_login(self):
        account = self.users["school_admin"]
        self.school.status = "suspended"
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        response = client.post(
            "/login?portal=admin",
            data={
                "_csrf": "test-csrf",
                "portal": "admin",
                "username": account.username,
                "password": "Test@12345",
            },
            follow_redirects=True,
        )
        self.assertIn("school portal is currently unavailable", response.get_data(as_text=True).lower())
        self.school.status = "active"
        run.db.session.commit()

    def test_import_staging_does_not_store_plaintext_password(self):
        client = run.app.test_client()
        admin = self.users["school_admin"]
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
            session["_csrf"] = "test-csrf"
        raw_password = "StrongImport123"
        csv_data = (
            f"full_name,username,email,password\nImported Parent,imported_parent,parent2@example.com,{raw_password}\n"
        ).encode()
        response = client.post(
            "/admin/import",
            data={
                "_csrf": "test-csrf",
                "record_type": "parent",
                "file": (BytesIO(csv_data), "parents.csv"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 302)
        Job = run.app.feature_models["BulkImportJob"]
        job = Job.query.order_by(Job.id.desc()).first()
        self.assertNotIn(raw_password, job.staged_json)
        self.assertIn("_password_hash", job.staged_json)

    def test_role_login_redirects_to_the_correct_portal(self):
        expected = {
            "system_admin": "/owner",
            "school_admin": "/dashboard",
            "teacher": "/dashboard",
            "student": "/dashboard",
            "parent": "/parent",
            "accountant": "/finance",
        }
        accounts = {**self.users, "parent": self.parent}
        for role, target in expected.items():
            with self.subTest(role=role):
                client = run.app.test_client()
                with client.session_transaction() as session:
                    session["_csrf"] = "test-csrf"
                portal = {
                    "teacher": "teacher",
                    "student": "student",
                    "parent": "parent",
                }.get(role, "admin")
                response = client.post(
                    f"/login?portal={portal}",
                    data={
                        "_csrf": "test-csrf",
                        "portal": portal,
                        "username": accounts[role].username,
                        "password": "Test@12345",
                    },
                )
                self.assertEqual(response.status_code, 302)
                self.assertTrue(response.location.endswith(target), response.location)

    def test_cross_school_people_finance_and_communications_are_isolated(self):
        other_school = run.School(
            name="Private Tenant School",
            slug="private-tenant-school",
            academic_year=self.school.academic_year,
            term=self.school.term,
            onboarded=True,
        )
        run.db.session.add(other_school)
        run.db.session.flush()
        other_admin = run.User(
            school_id=other_school.id,
            role="school_admin",
            full_name="Private Tenant Admin",
            username="private_tenant_admin",
            password_hash=generate_password_hash("Test@12345"),
            must_change_password=False,
        )
        other_teacher = run.User(
            school_id=other_school.id,
            role="teacher",
            full_name="Secret Teacher Name",
            username="secret_teacher",
            password_hash=generate_password_hash("Test@12345"),
            must_change_password=False,
        )
        other_student_user = run.User(
            school_id=other_school.id,
            role="student",
            full_name="Secret Student Name",
            username="secret_student",
            password_hash=generate_password_hash("Test@12345"),
            must_change_password=False,
        )
        other_parent = run.User(
            school_id=other_school.id,
            role="parent",
            full_name="Secret Parent Name",
            username="secret_parent",
            password_hash=generate_password_hash("Test@12345"),
            must_change_password=False,
        )
        run.db.session.add_all([other_admin, other_teacher, other_student_user, other_parent])
        run.db.session.flush()
        other_student = run.Student(
            school_id=other_school.id,
            user_id=other_student_user.id,
            admission_no="SECRET-001",
        )
        run.db.session.add(other_student)
        run.db.session.flush()
        run.db.session.add(
            run.ParentStudent(
                school_id=other_school.id,
                parent_id=other_parent.id,
                student_id=other_student.id,
            )
        )
        payment = run.ParentReportPayment(
            school_id=other_school.id,
            parent_id=other_parent.id,
            student_id=other_student.id,
            academic_year=other_school.academic_year,
            term=other_school.term,
            reference="PRIVATE-TENANT-PAYMENT",
            amount_subunit=1000,
            currency="GHS",
            status="success",
        )
        communication = run.Communication(
            school_id=other_school.id,
            channel="email",
            audience="parent",
            recipient="secret@example.com",
            subject="PRIVATE TENANT MESSAGE",
            message="Private tenant communication",
            created_by=other_admin.id,
        )
        Invoice = run.app.platform_models["Invoice"]
        invoice = Invoice(
            school_id=other_school.id,
            student_id=other_student.id,
            invoice_number="PRIVATE-INVOICE",
            academic_year=other_school.academic_year,
            term=other_school.term,
            amount_subunit=5000,
            balance_subunit=5000,
        )
        run.db.session.add_all([payment, communication, invoice])
        run.db.session.commit()

        client = run.app.test_client()
        admin = self.users["school_admin"]
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
            session["_csrf"] = "test-csrf"

        for url, forbidden_text in [
            ("/students", "Secret Student Name"),
            ("/teachers", "Secret Teacher Name"),
            ("/parent-links", "Secret Parent Name"),
            ("/finance", "PRIVATE-INVOICE"),
            ("/communications", "PRIVATE TENANT MESSAGE"),
        ]:
            with self.subTest(url=url):
                response = client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertNotIn(forbidden_text, response.get_data(as_text=True))

        response = client.post(
            "/finance",
            data={
                "_csrf": "test-csrf",
                "action": "manual_payment",
                "invoice_id": invoice.id,
                "amount": "1.00",
                "reason": "Cross-school attempt",
            },
        )
        self.assertEqual(response.status_code, 403)
        run.db.session.refresh(invoice)
        self.assertEqual(invoice.balance_subunit, 5000)

    def test_cross_school_branding_file_cannot_be_downloaded(self):
        filename = "cross-school-branding-test.png"
        other_school = run.School(name="Branding Tenant", slug="branding-tenant", crest=filename)
        run.db.session.add(other_school)
        run.db.session.commit()
        path = run.UPLOAD_DIR / filename
        path.write_bytes(b"tenant-private-file")
        self.addCleanup(lambda: path.unlink(missing_ok=True))

        client = run.app.test_client()
        admin = self.users["school_admin"]
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
        response = client.get(f"/uploads/{filename}")
        self.assertEqual(response.status_code, 404)

    def test_school_portal_login_uses_tenant_branding(self):
        self.school.name = "Branded Academy"
        self.school.slug = "branded-academy"
        self.school.motto = "Knowledge and Service"
        self.school.primary_color = "#123456"
        self.school.secondary_color = "#345678"
        self.school.accent_color = "#198754"
        run.db.session.commit()

        response = run.app.test_client().get(f"/school/{self.school.slug}/login?portal=admin")
        page = response.get_data(as_text=True)

        self.assertEqual(response.status_code, 200)
        self.assertIn("Branded Academy", page)
        self.assertIn("Knowledge and Service", page)
        self.assertIn("#345678", page)
        self.assertIn(f'name="school" value="{self.school.slug}"', page)

    def test_school_portal_rejects_credentials_from_another_tenant(self):
        other_school = run.School(name="Second Portal", slug="second-portal", onboarded=True)
        run.db.session.add(other_school)
        run.db.session.flush()
        other_admin = run.User(
            school_id=other_school.id,
            role="school_admin",
            full_name="Second Portal Admin",
            username=self.users["school_admin"].username,
            password_hash=generate_password_hash("Different@Test123"),
            must_change_password=False,
        )
        run.db.session.add(other_admin)
        run.db.session.commit()

        client = run.app.test_client()
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        response = client.post(
            f"/school/{other_school.slug}/login?portal=admin",
            data={
                "_csrf": "test-csrf",
                "portal": "admin",
                "school": other_school.slug,
                "username": self.users["school_admin"].username,
                "password": "Test@12345",
            },
            follow_redirects=True,
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn("Invalid username, password, or portal", response.get_data(as_text=True))

    def test_school_portal_serves_only_its_configured_public_favicon(self):
        filename = "tenant-public-favicon.png"
        self.school.slug = "favicon-school"
        self.school.favicon = filename
        run.db.session.commit()
        path = run.UPLOAD_DIR / filename
        path.write_bytes(b"favicon-content")
        self.addCleanup(lambda: path.unlink(missing_ok=True))

        client = run.app.test_client()
        page = client.get("/school/favicon-school/login")
        asset = client.get("/school/favicon-school/branding/favicon")

        self.assertEqual(page.status_code, 200)
        self.assertIn("/school/favicon-school/branding/favicon", page.get_data(as_text=True))
        self.assertEqual(asset.status_code, 200)
        self.assertEqual(asset.data, b"favicon-content")
        self.assertEqual(client.get("/school/favicon-school/branding/background").status_code, 404)

    def test_payment_receipt_is_branded_and_tenant_scoped(self):
        Invoice = run.app.platform_models["Invoice"]
        FinancePayment = run.app.platform_models["FinancePayment"]
        invoice = Invoice(
            school_id=self.school.id,
            student_id=self.student.id,
            invoice_number="RECEIPT-TEST-INVOICE",
            academic_year=self.school.academic_year,
            term=self.school.term,
            amount_subunit=2500,
            balance_subunit=0,
            status="paid",
        )
        run.db.session.add(invoice)
        run.db.session.flush()
        payment = FinancePayment(
            school_id=self.school.id,
            invoice_id=invoice.id,
            student_id=self.student.id,
            reference="RECEIPT-TEST-PAYMENT",
            amount_subunit=2500,
            currency="GHS",
            reason="Tuition",
            recorded_by=self.users["school_admin"].id,
        )
        run.db.session.add(payment)
        run.db.session.commit()

        client = run.app.test_client()
        admin = self.users["school_admin"]
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
        response = client.get(f"/finance/payments/{payment.id}/receipt")

        self.assertEqual(response.status_code, 200)
        self.assertIn(self.school.name, response.get_data(as_text=True))
        self.assertIn("RECEIPT-TEST-PAYMENT", response.get_data(as_text=True))

    def test_new_school_registration_creates_first_admin_and_login_slip(self):
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["_csrf"] = "test-csrf"
        response = client.post(
            "/register-school",
            data={
                "_csrf": "test-csrf",
                "school_name": "Onboarding Academy",
                "slug": "onboarding-academy",
                "admin_name": "First Administrator",
                "username": "onboarding_admin",
                "password": "Strong@Test123",
                "email": "onboarding@example.com",
                "phone": "0200000000",
            },
        )

        school = run.School.query.filter_by(slug="onboarding-academy").first()
        admin = run.User.query.filter_by(school_id=school.id, role="school_admin").first()
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith("/login-slip"))
        self.assertFalse(school.onboarded)
        self.assertTrue(admin.must_change_password)

    def test_unfinished_school_resumes_saved_setup_step(self):
        school = run.School(name="Resume School", slug="resume-school", onboarded=False, onboarding_step=3)
        run.db.session.add(school)
        run.db.session.flush()
        admin = run.User(
            school_id=school.id,
            role="school_admin",
            full_name="Resume Admin",
            username="resume_admin",
            password_hash=generate_password_hash("Strong@Test123"),
            must_change_password=False,
        )
        run.db.session.add(admin)
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version

        response = client.get("/dashboard")

        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.location.endswith("/setup/3"))

    def test_onboarding_cannot_launch_without_required_branding(self):
        school = run.School(
            name="Incomplete Launch School",
            slug="incomplete-launch",
            academic_year="2026/2027",
            term="Term 1",
            onboarded=False,
            onboarding_step=14,
        )
        run.db.session.add(school)
        run.db.session.flush()
        admin = run.User(
            school_id=school.id,
            role="school_admin",
            full_name="Launch Admin",
            username="launch_admin",
            password_hash=generate_password_hash("Strong@Test123"),
            must_change_password=False,
        )
        run.db.session.add(admin)
        run.db.session.commit()
        client = run.app.test_client()
        with client.session_transaction() as session:
            session["user_id"] = admin.id
            session["session_version"] = admin.session_version
            session["_csrf"] = "test-csrf"

        response = client.post("/setup/14", data={"_csrf": "test-csrf", "notes": "Ready"})
        run.db.session.refresh(school)

        self.assertEqual(response.status_code, 200)
        self.assertFalse(school.onboarded)
        self.assertIn("School logo", response.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
